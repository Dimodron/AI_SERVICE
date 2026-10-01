"""Deterministic, bounded report calculations. No model-supplied code or SQL."""
import csv
import io
import re
from decimal import Decimal, InvalidOperation
from uuid import uuid4

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from starlette.concurrency import run_in_threadpool

from config import settings
from schemas.report_template import ReportTemplate, validate_template_sources
from schemas.scenario_query import ScenarioQuery, QueryFilter
from schemas.reports import ReportResponse
from services.reports import MEDIA_TYPES, _text, _csv_cell

SCALES = {"rubles": 1, "thousands": 1000, "millions": 1000000, "billions": 1000000000}


def numeric(value):
    if value is None:
        return None
    try:
        result = Decimal(str(value))
        if not result.is_finite():
            raise ValueError("Нечисловое или бесконечное значение")
        return result
    except (InvalidOperation, ValueError):
        raise ValueError(f"Ожидалось число, получено {str(value)[:80]}") from None


def derived(column, values):
    left, right = values[column.left], values[column.right]
    if left is None or right is None:
        return None
    if column.operation == "difference":
        return numeric(left) - numeric(right)
    denominator = numeric(right)
    return numeric(left) / denominator if denominator else None


def matches(value, choices):
    if value is None:
        return False
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        try:
            return any(numeric(value) == numeric(choice) for choice in choices)
        except ValueError:
            raise ValueError("Условие числового поля содержит нечисловое значение") from None
    return str(value) in choices


def calculate(template, rows):
    dimensions = [c for c in template.columns if c.operation == "field"]
    aggregated = any(c.operation in ("sum", "count") for c in template.columns)
    groups = {}
    if aggregated:
        grouping = list(dict.fromkeys([*template.group_by, *(c.field for c in dimensions)]))
        for row in rows:
            key = tuple(row[field] for field in grouping)
            groups.setdefault(key, []).append(row)
        if not dimensions and not groups:
            groups[()] = []
        batches = list(groups.values())
    else:
        batches = [[row] for row in rows]
    output = []
    seen = set()
    for batch in batches:
        values = {}
        for col in template.columns:
            if col.operation == "field":
                value = batch[0][col.field]
                values[col.key] = numeric(value) if col.display != "text" else value
            elif col.operation in ("sum", "count"):
                subset = batch
                if col.condition_column:
                    subset = [r for r in batch if matches(r[col.condition_column], col.condition_values)]
                raw = [r[col.field] for r in subset] if col.field else [1 for _ in subset]
                if col.operation == "count":
                    values[col.key] = Decimal(sum(x is not None for x in raw))
                else:
                    numbers = [numeric(x) for x in raw if x is not None]
                    values[col.key] = sum(numbers, Decimal(0)) if numbers else (None if raw else Decimal(0))
            elif col.operation == "row_number":
                values[col.key] = None
            else:
                values[col.key] = derived(col, values)
        fingerprint = tuple(str(values[c.key]) if values[c.key] is not None else None for c in template.columns if c.operation != "row_number")
        if not template.distinct or fingerprint not in seen:
            output.append(values)
            seen.add(fingerprint)
    # Stable order even if PostgreSQL scan order changes.
    output.sort(key=lambda r: tuple(str(r[c.key]) if r[c.key] is not None else "" for c in template.columns))
    if template.sort_by:
        key = template.sort_by
        present = [r for r in output if r[key] is not None]
        missing = [r for r in output if r[key] is None]
        present.sort(key=lambda r: r[key], reverse=template.descending)
        output = present + missing
    for index, row in enumerate(output, 1):
        for col in template.columns:
            if col.operation == "row_number":
                row[col.key] = index
    totals = {}
    if template.totals:
        for col in template.columns:
            if col.operation in ("difference", "ratio", "percent"):
                totals[col.key] = derived(col, totals)
            elif col.display != "text" and col.operation != "row_number":
                numbers = [numeric(row[col.key]) for row in output if row[col.key] is not None]
                totals[col.key] = sum(numbers, Decimal(0)) if numbers else None
            else:
                totals[col.key] = None
    return output, totals


def display_value(template, col, value, csv_mode=False):
    if value is None:
        return "—"
    if col.operation == "row_number":
        return value
    if col.display == "text":
        return _text(value)
    value = numeric(value)
    if col.money:
        value /= SCALES[template.units]
    if csv_mode:
        if col.display == "percent":
            return f"{value * 100:.{col.decimals}f}%"
        return f"{value:.{col.decimals}f}"
    return value


def column_title(template, column):
    units = {"rubles": "руб.", "thousands": "тыс. руб.", "millions": "млн руб.", "billions": "млрд руб."}
    return column.title + (" (" + units[template.units] + ")" if column.money else "")


def render(template, rows, totals, filters):
    headings = [column_title(template, c) for c in template.columns]
    values = [[display_value(template, col, row[col.key], template.format == "csv") for col in template.columns] for row in rows]
    if template.totals:
        total_row = [display_value(template, c, totals[c.key], template.format == "csv") if c.total else "" for c in template.columns]
        label = next((i for i, c in enumerate(template.columns) if not c.total), None)
        if label is not None:
            total_row[label] = "ИТОГО"
        values.append(total_row)
    if template.format == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        # CSV remains a rectangular data table: no title/filter preamble.
        writer.writerow(headings)
        writer.writerows([_csv_cell(x) for x in row] for row in values)
        return stream.getvalue().encode("utf-8-sig")
    book = Workbook()
    sheet = book.active
    sheet.title = template.sheet_name
    def append(row):
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.value = _text(cell.value)
                cell.data_type = "s"
            if isinstance(cell.value, int) and abs(cell.value) >= 10**15:
                cell.value = str(cell.value)
                cell.data_type = "s"
    append([template.title])
    if template.show_filters:
        append(["Фильтры: " + filters])
    if template.units != "rubles":
        append(["Единицы сумм: " + {"thousands": "тыс. руб.", "millions": "млн руб.", "billions": "млрд руб."}[template.units]])
    append(headings)
    header = sheet.max_row
    for cell in sheet[header]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="304B63")
    for row in values:
        append(row)
    for i, col in enumerate(template.columns, 1):
        sheet.column_dimensions[get_column_letter(i)].width = col.width
        fmt = "0" if col.operation == "row_number" else ("#,##0" + ("." + "0"*col.decimals if col.decimals else "") + ("%" if col.display == "percent" else ""))
        for cells in sheet.iter_rows(min_row=header+1, min_col=i, max_col=i):
            cell = cells[0]
            cell.alignment = Alignment(wrap_text=template.wrap_text, vertical="top")
            if col.display != "text" or col.operation == "row_number":
                cell.number_format = fmt
    if template.freeze_header:
        sheet.freeze_panes = f"A{header+1}"
    if template.autofilter:
        sheet.auto_filter.ref = f"A{header}:{get_column_letter(len(headings))}{header+len(rows)}"
    if template.totals:
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


async def create_template_report(connection, scenario, request, payload, *, is_admin=False):
    # Import locally to share all query authorization checks without an import cycle.
    from services.scenario_runner import query_scenario
    if not scenario.get("is_active", True):
        raise ValueError("Сценарий выключен")
    if not scenario.get("report_template"):
        raise ValueError("У сценария нет шаблона отчёта")
    template = ReportTemplate.model_validate(scenario["report_template"])
    validate_template_sources(template, scenario.get("tables") or [], scenario.get("table_name"), scenario.get("columns_description"))
    if request.format:
        template = template.model_copy(update={"format": request.format})
    params = {p.column: p for p in template.parameters}
    if set(request.parameters) - set(params):
        raise ValueError("Неизвестные параметры шаблона: " + ", ".join(sorted(set(request.parameters)-set(params))))
    filters = list(template.filters)
    for column, param in params.items():
        value = request.parameters.get(column, param.default)
        if value is None or not value.strip():
            if param.required:
                raise ValueError("Не указан обязательный параметр: " + param.title)
            continue
        if len(value) > 500:
            raise ValueError("Значение параметра слишком длинное")
        filters.append(QueryFilter(column=column, value=value))
    fields = set(template.group_by)
    for col in template.columns:
        if col.field:
            fields.add(col.field)
        if col.condition_column:
            fields.add(col.condition_column)
    if not fields:
        raise ValueError("Добавьте хотя бы одно поле источника")
    query = ScenarioQuery(scenario_id=request.scenario_id, table_name=template.table_name,
                          columns=sorted(fields), filters=filters)
    data = await query_scenario(connection, scenario, query, payload, is_admin=is_admin,
                                export_limit=settings.MAX_TEMPLATE_ROWS)
    if data["truncated"]:
        raise ValueError(f"Выборка превышает {settings.MAX_TEMPLATE_ROWS} строк. Уточните фильтры; неполный файл не создан.")
    rows, totals = await run_in_threadpool(calculate, template, data["rows"])
    descriptions = [f"{f.column} {f.operator} {f.value}" for f in filters]
    if not is_admin:
        descriptions.append(f"jurpers = {payload.user_jurpers}")
        if payload.user_organization is not None:
            descriptions.append(f"organization = {payload.user_organization}")
    filter_text = "; ".join(descriptions) or "Без дополнительных фильтров"
    content = await run_in_threadpool(render, template, rows, totals, filter_text)
    if len(content) > settings.MAX_TEMPLATE_BYTES:
        raise ValueError("Файл превышает MAX_TEMPLATE_BYTES; уточните фильтры")
    filename = (re.sub(r"[^\w .-]", "_", template.title).strip(" .") or "report") + "." + template.format
    file_id = uuid4()
    extracted = template.title + "\n" + "\t".join(column_title(template, c) for c in template.columns) + "\n" + "\n".join(
        "\t".join(str(display_value(template, c, r[c.key], True)) for c in template.columns) for r in rows)
    cur = await connection.execute(
        "INSERT INTO files(id,filename,media_type,size_bytes,extracted_text,content) VALUES (%s,%s,%s,%s,%s,%s) "
        "RETURNING id AS file_id,filename,media_type,size_bytes,created_at",
        (file_id, filename, MEDIA_TYPES[template.format], len(content), extracted, content))
    report = ReportResponse(**await cur.fetchone(), download_url=f"/api/files/{file_id}/download")
    def escape(value):
        return str(value).replace("|", "\\|").replace("\n", " ")
    preview = [" | ".join(escape(column_title(template, c)) for c in template.columns),
               " | ".join("---" for _ in template.columns)]
    for row in rows[:10]:
        preview.append(" | ".join(escape(display_value(template, c, row[c.key], True)) for c in template.columns))
    summary = f"Сформирован файл «{filename}». Строк данных: {len(rows)}.\n\n" + "\n".join(preview)
    if len(rows) > 10:
        summary += "\n\nПоказаны первые 10 строк. Полная выборка — в файле."
    if not rows:
        summary += "\n\nПо заданным фильтрам данных нет."
    return report, summary
