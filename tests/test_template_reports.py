import csv
import io
import os
import json
from decimal import Decimal
from pathlib import Path
from unittest import TestCase, IsolatedAsyncioTestCase, skipUnless
from unittest.mock import patch, AsyncMock
from uuid import uuid4

from openpyxl import load_workbook
from pydantic import ValidationError
from fastapi import HTTPException
from database.history import connect, database_lifespan
from schemas.report_template import ReportTemplate, TemplateReportRequest
from schemas.scenarios import ScenarioCreate, ScenarioUpdate
from schemas.qwen import ChatRequest
from services.template_reports import calculate, render, create_template_report
from services.scenarios import create_scenario, update_scenario
from services.scenario_runner import answer_with_scenarios


def expenses():
    return ReportTemplate(title="Расходы", table_name="oracle_data.template_source", units="millions",
                          group_by=["organization"], totals=True, sort_by="total", descending=True,
                          columns=[
        {"key":"number","title":"№","operation":"row_number"},
        {"key":"name","title":"Учреждение","field":"orgname","width":60},
        {"key":"total","title":"Всего","operation":"sum","field":"esum","display":"number","total":True,"money":True},
        {"key":"salary","title":"Зарплата","operation":"sum","field":"esum","condition_column":"kvrcode",
         "condition_values":["111","119"],"display":"number","total":True,"money":True},
        {"key":"share","title":"Доля зарплаты","operation":"percent","left":"salary","right":"total",
         "display":"percent","decimals":1,"total":True},
    ])


class CalculationTests(TestCase):
    def test_grouping_numeric_totals_and_formatting(self):
        template=expenses()
        rows=[
            {"organization":1,"orgname":"A","esum":Decimal("1000000"),"kvrcode":Decimal("111.0")},
            {"organization":1,"orgname":"A","esum":Decimal("3000000"),"kvrcode":"244"},
            {"organization":2,"orgname":"B","esum":Decimal("2000000"),"kvrcode":"119"},
            {"organization":3,"orgname":"C","esum":Decimal("0"),"kvrcode":"111"},
        ]
        data, totals=calculate(template,rows)
        self.assertEqual([r["number"] for r in data],[1,2,3])
        self.assertEqual(data[0]["share"],Decimal("0.25"))
        self.assertIsNone(data[2]["share"])
        self.assertEqual(totals["share"],Decimal("0.5"))
        binary=render(template,data,totals,"version = 123")
        book=load_workbook(io.BytesIO(binary))
        sheet=book.active
        self.assertEqual(sheet.freeze_panes,"A5")
        self.assertEqual(sheet.cell(5,3).value,4)
        self.assertEqual(sheet.cell(5,5).value,0.25)
        self.assertEqual(sheet.cell(5,5).number_format,"#,##0.0%")
        self.assertEqual(sheet.cell(8,5).value,0.5)
        self.assertEqual(sheet.auto_filter.ref,"A4:E7")
        self.assertEqual(sheet.column_dimensions["B"].width,60)
        book.close()
        csv_template=template.model_copy(update={"format":"csv"})
        parsed=list(csv.reader(io.StringIO(render(csv_template,data,totals,"").decode("utf-8-sig"))))
        self.assertEqual(parsed[1][2],"4.00")
        self.assertEqual(parsed[-1][-1],"50.0%")
        self.assertTrue(all(len(row)==5 for row in parsed))

    def test_count_difference_ratio_and_null_denominator(self):
        template=ReportTemplate(title="Stats",table_name="data",totals=True,units="millions",columns=[
            {"key":"total","title":"Total","operation":"sum","field":"value","display":"number","total":True},
            {"key":"n","title":"Count","operation":"count","display":"number","total":True},
            {"key":"difference","title":"Difference","operation":"difference","left":"total","right":"n","display":"number","total":True},
            {"key":"ratio","title":"Ratio","operation":"ratio","left":"total","right":"n","display":"number","total":True},
        ])
        rows, totals=calculate(template,[{"value":Decimal(4)},{"value":None},{"value":Decimal(8)}])
        self.assertEqual(rows[0]["n"],3)
        self.assertEqual(rows[0]["difference"],9)
        self.assertEqual(rows[0]["ratio"],4)
        self.assertEqual(totals["ratio"],4)
        from services.template_reports import display_value
        self.assertEqual(display_value(template,template.columns[0],rows[0]["total"]),12) # not a money column
        empty,_=calculate(template,[])
        self.assertIsNone(empty[0]["ratio"])

    def test_duplicate_names_remain_separate_by_id(self):
        t=expenses()
        data,_=calculate(t,[{"organization":i,"orgname":"Same","esum":Decimal(1),"kvrcode":"111"} for i in (1,2)])
        self.assertEqual(len(data),2)

    def test_schema_rejects_forward_refs_and_unavailable_fields(self):
        broken=expenses().model_dump()
        broken["columns"][4]["left"]="future"
        with self.assertRaises(ValidationError):ReportTemplate.model_validate(broken)
        with self.assertRaises(ValidationError):
            ScenarioCreate(title="Bad",scenario="test",tables=[{"table_name":"oracle_data.template_source","columns_description":{"orgname":"Name"}}],report_template=expenses())

    def test_formula_like_text_is_not_executed(self):
        t=ReportTemplate(title="Safe",table_name="example",columns=[{"key":"name","title":"Name","field":"name"}])
        data,totals=calculate(t,[{"name":"=SUM(1,2)"}])
        book=load_workbook(io.BytesIO(render(t,data,totals,"")))
        self.assertEqual(book.active.cell(4,1).data_type,"s")
        book.close()
        raw=render(t.model_copy(update={"format":"csv"}),data,totals,"").decode("utf-8-sig")
        self.assertIn("'=SUM",raw)


@skipUnless(os.getenv("TEST_DATABASE")=="1","requires disposable PostgreSQL")
class TemplateDatabaseTests(IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.life=database_lifespan();await self.life.__aenter__()
        async with await connect() as conn:
            await conn.execute((Path(__file__).resolve().parents[1]/"database/init.sql").read_text())
            await conn.execute("CREATE SCHEMA IF NOT EXISTS oracle_data")
            await conn.execute("DROP TABLE IF EXISTS oracle_data.template_source")
            await conn.execute("CREATE TABLE oracle_data.template_source(jur_pers bigint,organization bigint,orgname text,esum numeric,kvrcode text,version bigint)")
            await conn.execute("INSERT INTO oracle_data.template_source SELECT 10,i,'Org '||i,100,'111',123 FROM generate_series(1,444) i")
            await conn.execute("INSERT INTO oracle_data.template_source VALUES(99,999,'Hidden',9999,'111',123)")
        self.payload=ChatRequest(user_login="alice",user_jurpers=10,message="Дай Excel")
        t=ReportTemplate(title="Учреждения",table_name="oracle_data.template_source",distinct=True,
                         parameters=[{"column":"version","title":"Версия"}],
                         columns=[{"key":"n","title":"№","operation":"row_number"},{"key":"name","title":"Учреждение","field":"orgname"}])
        self.scenario={"id":uuid4(),"tables":[{"table_name":"oracle_data.template_source","columns_description":{k:k for k in ["jur_pers","organization","orgname","esum","kvrcode","version"]}}],
                       "is_admin":False,"visible_jurpers":[],"is_active":True,"report_template":t.model_dump()}
        self.request=TemplateReportRequest(scenario_id=self.scenario["id"],parameters={"version":"123"})

    async def asyncTearDown(self):
        await self.life.__aexit__(None,None,None)

    async def test_full_export_scope_required_filters_and_row_limit(self):
        async with await connect() as conn:
            report,summary=await create_template_report(conn,self.scenario,self.request,self.payload)
            binary=(await (await conn.execute("SELECT content FROM files WHERE id=%s",(report.file_id,))).fetchone())["content"]
            book=load_workbook(io.BytesIO(binary)); sheet=book.active
            self.assertEqual(sheet.max_row,447) # title, filters, header + 444 rows
            self.assertNotIn("Hidden",[sheet.cell(r,2).value for r in range(4,448)])
            self.assertEqual(sheet.cell(447,1).value,444)
            book.close()
            self.assertIn("444",summary)
            with self.assertRaises(ValueError):
                await create_template_report(conn,self.scenario,self.request.model_copy(update={"parameters":{}}),self.payload)
            with patch("services.template_reports.settings", type("Settings",(),{"MAX_TEMPLATE_ROWS":200})()):
                with self.assertRaisesRegex(ValueError,"неполный файл"):
                    await create_template_report(conn,self.scenario,self.request,self.payload)
            with self.assertRaises(ValueError):
                await create_template_report(conn,{**self.scenario,"is_admin":True},self.request,self.payload)

    async def test_template_crud_preserves_and_validates_merged_source(self):
        row=await create_scenario(ScenarioCreate(title="Template",scenario="Export",tables=self.scenario["tables"],report_template=self.scenario["report_template"]))
        self.assertEqual(row["report_template"]["title"],"Учреждения")
        updated=await update_scenario(row["id"],ScenarioUpdate(description="Changed"))
        self.assertEqual(updated["report_template"],row["report_template"])
        with self.assertRaises(HTTPException) as error:
            await update_scenario(row["id"],ScenarioUpdate(tables=[]))
        self.assertEqual(error.exception.status_code,422)
        removed=await update_scenario(row["id"],ScenarioUpdate(report_template=None,tables=[]))
        self.assertIsNone(removed["report_template"])

    async def test_model_cannot_complete_explicit_file_request_without_file(self):
        request_message=json.dumps({"status":"completed","response":"Файл уже готов"})
        reply=AsyncMock(side_effect=[
            {"role":"assistant","content":request_message},
            {"role":"assistant","content":"","tool_calls":[{"function":{"name":"create_template_report","arguments":self.request.model_dump(mode="json")}}]}
        ])
        async with await connect() as conn:
            with patch("services.scenario_runner.QwenStrategy.chat_message",reply):
                answer,files=await answer_with_scenarios(conn,"test",self.payload,[],{str(self.scenario["id"]):self.scenario})
        self.assertEqual(reply.await_count,2)
        self.assertEqual(len(files),1)
        self.assertIn(files[0].download_url,answer)

    async def test_tool_returns_server_file_without_model_rewriting(self):
        reply=AsyncMock(return_value={"role":"assistant","content":"","tool_calls":[{"function":{"name":"create_template_report","arguments":self.request.model_dump(mode="json")}}]})
        async with await connect() as conn:
            with patch("services.scenario_runner.QwenStrategy.chat_message",reply):
                answer,files=await answer_with_scenarios(conn,"test",self.payload,[],{str(self.scenario["id"]):self.scenario})
        self.assertEqual(reply.await_count,1)
        self.assertEqual(len(files),1)
        self.assertIn(files[0].download_url,answer)
        self.assertIn("444",answer)
